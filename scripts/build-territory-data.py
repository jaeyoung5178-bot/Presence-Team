#!/usr/bin/env python3
"""Build the compact browser territory dataset from the official workbook.

The source workbook remains outside the repository. Exact duplicate rows are
removed in source order; distinct codes for the same site are preserved.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


SHEET = "2026년 스트릿 사이트명&테리토리 코드"


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\r", " ").replace("\n", " ").replace("\t", " ")


def build(source: Path) -> tuple[str, str, dict[str, object]]:
    workbook = load_workbook(source, read_only=False, data_only=True)
    sheet = workbook[SHEET]

    rows: list[tuple[str, str, str, str]] = []
    seen_rows: set[tuple[str, str, str, str]] = set()
    duplicate_rows = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
        sido, sigu, name, code = [clean(value) for value in values[:4]]
        if not any((sido, sigu, name, code)):
            continue
        if not (sido and name and code):
            raise ValueError(f"row {row_number}: required territory value is blank")
        record = (sido, sigu, name, code)
        if record in seen_rows:
            duplicate_rows += 1
            continue
        seen_rows.add(record)
        rows.append(record)

    locations: list[tuple[str, str]] = []
    location_index: dict[tuple[str, str], int] = {}
    compact_rows: list[str] = []
    for sido, sigu, name, code in rows:
        location = (sido, sigu)
        if location not in location_index:
            location_index[location] = len(locations)
            locations.append(location)
        compact_rows.append(f"{location_index[location]}\t{name}\t{code}")

    loc_blob = "\x1e".join(f"{sido}\x1f{sigu}" for sido, sigu in locations)
    row_blob = "\n".join(compact_rows)
    meta = {
        "source": source.name,
        "sheet": SHEET,
        "records": len(rows),
        "locations": len(locations),
        "exactDuplicatesRemoved": duplicate_rows,
        "version": "2026-10-25",
    }
    return loc_blob, row_blob, meta


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    loc_blob, row_blob, meta = build(args.source)
    content = (
        "/* Generated from the official 2026 Street Territory workbook. */\n"
        f"window.__TERRI_2026_META={json.dumps(meta, ensure_ascii=False, separators=(',', ':'))};\n"
        f"window.__TERRI_2026_LOC={json.dumps(loc_blob, ensure_ascii=False)};\n"
        f"window.__TERRI_2026_ROWS={json.dumps(row_blob, ensure_ascii=False)};\n"
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(content, encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False))


if __name__ == "__main__":
    main()
